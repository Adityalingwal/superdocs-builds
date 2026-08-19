from pathlib import Path

SUPPORTED_SUFFIXES = (".md", ".txt", ".pdf", ".docx")


def read_document(path: Path) -> str:
    """Read one source document as text, whatever its supported format.

    A format outside the supported list, a damaged file, or a PDF with no
    extractable text is refused with the cause and the practical fix —
    never silently skipped or half-read.
    """
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(str(path))
            text = "\n".join(
                page.extract_text() or "" for page in reader.pages
            ).strip()
        except Exception as e:
            raise ValueError(
                f"could not read {path.name}: {e} — re-export the PDF "
                f"(unprotected, non-corrupt) or provide the text as .md/.txt"
            ) from e
        if not text:
            raise ValueError(
                f"{path.name} contains no extractable text — a scanned or "
                f"image-only PDF needs OCR, which this build does not do; "
                f"export a text-based PDF or provide the text as .md/.txt"
            )
        return text
    if suffix == ".docx":
        import docx

        try:
            document = docx.Document(str(path))
        except Exception as e:
            raise ValueError(
                f"could not read {path.name}: {e} — re-save the file as "
                f".docx from Word (a renamed .doc is not a .docx) or "
                f"provide the text as .md/.txt"
            ) from e
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if not text:
            raise ValueError(
                f"{path.name} contains no readable text — the document is "
                f"empty or its content is only images; provide the text as "
                f".md/.txt"
            )
        return text
    raise ValueError(
        f"{path.name}: format '{suffix}' is not supported — supported "
        f"formats are {', '.join(SUPPORTED_SUFFIXES)}"
    )
