from pathlib import Path

import pytest

from engine.read_document import read_document
from engine.requests_from_notice import parse_notice

FIXTURES = Path(__file__).parent / "fixtures"


def write_pdf(path: Path, text: str) -> Path:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=letter)
    y = letter[1] - 60
    for line in text.splitlines():
        if y < 60:
            pdf.showPage()
            y = letter[1] - 60
        pdf.drawString(50, y, line)
        y -= 14
    pdf.save()
    return path


@pytest.fixture(scope="module")
def pdf_notice(tmp_path_factory):
    text = (FIXTURES / "notice-style-b-numbered-list.md").read_text(encoding="utf-8")
    return write_pdf(tmp_path_factory.mktemp("pdf") / "rfe-notice.pdf", text)


def test_a_pdf_notice_parses_into_the_same_five_requests(pdf_notice):
    requests = parse_notice(read_document(pdf_notice))
    assert [r.id for r in requests] == ["R1", "R2", "R3", "R4", "R5"]
    assert "nonimmigrant" in requests[4].text


def test_markdown_and_text_files_read_unchanged(tmp_path):
    file = tmp_path / "note.txt"
    file.write_text("plain text content", encoding="utf-8")
    assert read_document(file) == "plain text content"


def test_a_textless_pdf_is_refused_naming_ocr_as_the_gap(tmp_path):
    from pypdf import PdfWriter

    blank = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with open(blank, "wb") as f:
        writer.write(f)
    with pytest.raises(ValueError, match="OCR"):
        read_document(blank)


def test_an_unsupported_format_is_refused_with_the_supported_list(tmp_path):
    file = tmp_path / "notice.xlsx"
    file.write_text("whatever", encoding="utf-8")
    with pytest.raises(ValueError, match="supported"):
        read_document(file)


def test_a_damaged_pdf_is_refused_with_the_cause_and_fix(tmp_path):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 garbage that is not a real pdf")
    with pytest.raises(ValueError, match="re-export"):
        read_document(broken)


def write_docx(path: Path, text: str) -> Path:
    import docx

    document = docx.Document()
    for line in text.splitlines():
        if line.strip():
            document.add_paragraph(line)
    document.save(str(path))
    return path


def test_a_docx_notice_parses_into_the_same_five_requests(tmp_path):
    text = (FIXTURES / "notice-style-b-numbered-list.md").read_text(encoding="utf-8")
    docx_notice = write_docx(tmp_path / "rfe-notice.docx", text)
    requests = parse_notice(read_document(docx_notice))
    assert [r.id for r in requests] == ["R1", "R2", "R3", "R4", "R5"]
    assert "nonimmigrant" in requests[4].text


def test_docx_table_text_is_readable_too(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Exhibit summary")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Exhibit C"
    table.rows[0].cells[1].text = "Degree certificate and transcripts"
    file = tmp_path / "petition-part.docx"
    document.save(str(file))
    text = read_document(file)
    assert "Exhibit summary" in text
    assert "Degree certificate and transcripts" in text


def test_a_damaged_docx_is_refused_with_the_cause_and_fix(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"PK garbage that is not a real docx")
    with pytest.raises(ValueError, match="re-export|re-save"):
        read_document(broken)
